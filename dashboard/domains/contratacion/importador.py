"""
Importador del libro maestro de contratación (`data/Contratos SRNI 2026.xlsx`).

Lógica pura reutilizada por:
  - management/commands/importar_contratos.py (cron / CLI)
  - domains/contratacion/views.py::contratacion_sincronizar (botón "Sincronizar")

Diseño: upsert idempotente. El maestro de contratos (`Contrato`) se actualiza por
`numero_contrato`; los pagos y la seguridad social se reemplazan por contrato en cada
corrida. Todo ocurre dentro de una transacción; con `dry_run=True` se hace rollback.
"""
import datetime
import unicodedata
from decimal import Decimal, InvalidOperation

from django.conf import settings
from django.db import transaction

from dashboard.domains.importacion.services import ImportacionValidationError
from dashboard.models import (
    Colaborador,
    Contrato,
    ContratoImportLog,
    Convenio,
    PagoContrato,
    SeguridadSocialContrato,
)

MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]

ESTADO_MAP = {
    "EJECUCION": "ejecucion",
    "CESION": "cesion",
    "TERMINADO": "terminado",
    "SUSPENDIDO": "suspendido",
}


def _cargar_workbook(ruta):
    """
    Carga el libro con `data_only=True`: muchas columnas (valor total, cancelado,
    saldo, % ejec, algunos pagos) son fórmulas y necesitamos su valor cacheado, no
    la fórmula. `read_only=True` acelera la lectura secuencial con iter_rows.
    """
    try:
        import openpyxl
        return openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    except Exception as exc:
        raise ImportacionValidationError(
            "No se pudo leer el archivo. Asegúrate de que sea un .xlsx válido."
        ) from exc


# --------------------------------------------------------------------------- #
# Helpers de normalización de celdas
# --------------------------------------------------------------------------- #
def _norm(texto):
    texto = str(texto).upper().strip()
    texto = " ".join(texto.split())  # colapsa espacios internos/dobles
    return "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )


def _header_map(header_row):
    return {
        _norm(valor): idx
        for idx, valor in enumerate(header_row)
        if valor is not None and str(valor).strip() != ""
    }


def _cell(row, hmap, nombre):
    idx = hmap.get(_norm(nombre))
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _to_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _to_cedula(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return str(int(v))
    s = str(v).strip()
    return s or None


def _to_decimal(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).replace("$", "").replace(",", "").strip()
    if not s:
        return None
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return None


def _estado(v):
    return ESTADO_MAP.get(_norm(v), "otro") if v else "otro"


def _numero_contrato(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        v = int(v)
    return str(v).strip() or None


# --------------------------------------------------------------------------- #
# Parseo de hojas auxiliares (SECOP II, SISEG)
# --------------------------------------------------------------------------- #
def _parse_secop(wb):
    """{numero_contrato: {mes(1-12): 'pagado'|'pendiente'}}"""
    resultado = {}
    if "SECOP II" not in wb.sheetnames:
        return resultado
    rows = list(wb["SECOP II"].iter_rows(values_only=True))
    if not rows:
        return resultado
    hmap = _header_map(rows[0])
    for row in rows[1:]:
        numero = _numero_contrato(_cell(row, hmap, "CONTRATO"))
        if not numero:
            continue
        estados = {}
        for i, mes_nombre in enumerate(MESES, start=1):
            valor = _cell(row, hmap, mes_nombre)
            if valor is not None:
                estados[i] = "pagado" if "PAGAD" in _norm(valor) else "pendiente"
        resultado[numero] = estados
    return resultado


def _parse_siseg(wb):
    """{numero_contrato: [{mes, codigo, fecha_aprobado, mes_cubierto}]}"""
    resultado = {}
    if "SISEG" not in wb.sheetnames:
        return resultado
    # En SISEG la fila de encabezados de columna es la 2 (la 1 es la agrupación por mes).
    rows = list(wb["SISEG"].iter_rows(values_only=True))
    if len(rows) < 2:
        return resultado
    # Localiza el índice de la columna CONTRATO y el inicio de los bloques CÓDIGO/APROBADO/SS.
    header = rows[1]
    contrato_idx = None
    bloque_inicio = None
    for idx, valor in enumerate(header):
        etiqueta = _norm(valor) if valor is not None else ""
        if etiqueta == "CONTRATO" and contrato_idx is None:
            contrato_idx = idx
        if etiqueta == "CODIGO" and bloque_inicio is None:
            bloque_inicio = idx
    if contrato_idx is None or bloque_inicio is None:
        return resultado

    for row in rows[2:]:
        if contrato_idx >= len(row):
            continue
        numero = _numero_contrato(row[contrato_idx])
        if not numero:
            continue
        registros = []
        mes = 1
        base = bloque_inicio
        while base + 2 < len(row) + 1 and mes <= 12:
            codigo = row[base] if base < len(row) else None
            aprobado = row[base + 1] if base + 1 < len(row) else None
            cubierto = row[base + 2] if base + 2 < len(row) else None
            if codigo or aprobado or cubierto:
                registros.append({
                    "mes": mes,
                    "codigo": _to_str(codigo),
                    "fecha_aprobado": _to_date(aprobado),
                    "mes_cubierto": _to_str(cubierto),
                })
            base += 3
            mes += 1
        if registros:
            resultado[numero] = registros
    return resultado


# --------------------------------------------------------------------------- #
# Procesamiento de contratos (Unidad + Proye. adiciones)
# --------------------------------------------------------------------------- #
def _campos_contrato(row, hmap, modalidad):
    porcentaje = _to_decimal(_cell(row, hmap, "%  EJEC."))
    if porcentaje is not None:
        porcentaje = (porcentaje * 100).quantize(Decimal("0.01"))
    return {
        "cedula": _to_cedula(_cell(row, hmap, "CEDULA")),
        "nombre_contratista": _to_str(_cell(row, hmap, "NOMBRE CONTRATISTA")) or "",
        "equipo": _to_str(_cell(row, hmap, "EQUIPO")),
        "id_externo": _to_str(_cell(row, hmap, "ID")),
        "estado": _estado(_cell(row, hmap, "ESTADO")),
        "modalidad": modalidad,
        "objeto": _to_str(_cell(row, hmap, "OBJETO")),
        "perfil_academico": _to_str(_cell(row, hmap, "PERFIL ACADÉMICO")),
        "cdp": _to_str(_cell(row, hmap, "CDP")),
        "rp": _to_str(_cell(row, hmap, "RP")),
        "fecha_inicio": _to_date(_cell(row, hmap, "INICIO"))
        or _to_date(_cell(row, hmap, "INICIO SECOP II")),
        "fecha_terminacion": _to_date(_cell(row, hmap, "TERMINACIÓN"))
        or _to_date(_cell(row, hmap, "TERMINACIÓN SECOP II")),
        "valor_honorarios": _to_decimal(_cell(row, hmap, "VALOR HONORARIOS")),
        "valor_total_contrato": _to_decimal(
            _cell(row, hmap, "VALOR REAL CONTRATO SEGÚN FECHA DE INICIO")
        ),
        "valor_cancelado": _to_decimal(_cell(row, hmap, "VALOR TOTAL CANCELADO")),
        "saldo_contrato": _to_decimal(_cell(row, hmap, "SALDO CONTRATO")),
        "porcentaje_ejecucion": porcentaje,
        "ibc": _to_decimal(_cell(row, hmap, "IBC")),
        "observaciones": _to_str(_cell(row, hmap, "OBSERVACIONES")),
    }


def _procesar_hoja_contratos(wb, hoja_nombre, modalidad, secop_map, siseg_map,
                             colaboradores, resumen):
    if hoja_nombre not in wb.sheetnames:
        return
    rows = list(wb[hoja_nombre].iter_rows(values_only=True))
    if not rows:
        return
    hmap = _header_map(rows[0])

    for numero_fila, row in enumerate(rows[1:], start=2):
        numero = _numero_contrato(_cell(row, hmap, "CONTRATO"))
        nombre = _to_str(_cell(row, hmap, "NOMBRE CONTRATISTA"))
        if not numero or not nombre:
            continue
        resumen["filas_leidas"] += 1
        try:
            campos = _campos_contrato(row, hmap, modalidad)
            campos["colaborador"] = colaboradores.get(campos["cedula"])
            if campos["colaborador"] is None:
                resumen["sin_match_colaborador"] += 1

            defaults = {k: v for k, v in campos.items() if k != "modalidad"}
            contrato, creado = Contrato.objects.update_or_create(
                numero_contrato=numero,
                modalidad=modalidad,
                defaults=defaults,
            )
            resumen["creados" if creado else "actualizados"] += 1

            # Pagos (PAGO 1–12) + estado SECOP II — reemplazo idempotente.
            estados_secop = secop_map.get(numero, {})
            contrato.pagos.all().delete()
            pagos = []
            for mes in range(1, 13):
                valor = _to_decimal(_cell(row, hmap, f"PAGO {mes}"))
                estado_secop = estados_secop.get(mes, "pendiente")
                if valor is None and mes not in estados_secop:
                    continue
                pagos.append(PagoContrato(
                    contrato=contrato,
                    mes=mes,
                    valor=valor,
                    estado_secop=estado_secop,
                ))
            if pagos:
                PagoContrato.objects.bulk_create(pagos)

            # Seguridad social (SISEG) — reemplazo idempotente.
            contrato.seguridad_social.all().delete()
            registros_ss = siseg_map.get(numero, [])
            if registros_ss:
                SeguridadSocialContrato.objects.bulk_create([
                    SeguridadSocialContrato(contrato=contrato, **reg)
                    for reg in registros_ss
                ])
        except Exception as exc:  # noqa: BLE001 - se registra por fila y se continúa
            resumen["errores"].append(f"{hoja_nombre} fila {numero_fila} ({numero}): {exc}")


def _procesar_convenios(wb, hoja_nombre, tipo, resumen):
    if hoja_nombre not in wb.sheetnames:
        return
    rows = list(wb[hoja_nombre].iter_rows(values_only=True))
    if not rows:
        return
    hmap = _header_map(rows[0])

    # Reemplazo idempotente por tipo.
    Convenio.objects.filter(tipo=tipo).delete()
    nuevos = []
    for row in rows[1:]:
        entidad = _to_str(_cell(row, hmap, "ENTIDAD"))
        if not entidad:
            continue
        fecha_term = _cell(row, hmap, "FECHA TERMINACIÓN")
        fecha_term = _to_str(fecha_term) if fecha_term is not None else None
        num_mod = _cell(row, hmap, "No. MODIFICATORIOS")
        nuevos.append(Convenio(
            entidad=entidad,
            articulador=_to_str(_cell(row, hmap, "ARTICULADOR")),
            seyco=_to_str(_cell(row, hmap, "SEYCO")),
            numero_convenio=_numero_contrato(_cell(row, hmap, "No. CONVENIO")),
            fecha_suscripcion=_to_date(_cell(row, hmap, "FECHA SUSCRIPCIÓN")),
            fecha_terminacion=fecha_term,
            num_modificatorios=str(int(num_mod)) if isinstance(num_mod, (int, float)) else _to_str(num_mod),
            contratista=_to_str(_cell(row, hmap, "CONTRATISTA")),
            clase_contrato=_to_str(_cell(row, hmap, "CLASE DE CONTRATO")),
            valor_inicial=_to_decimal(_cell(row, hmap, "VALOR INICIAL CONTRATO")),
            objeto=_to_str(_cell(row, hmap, "OBJETO")),
            link=_to_str(_cell(row, hmap, "LINK DEL ACUERDO O CONVENIO")),
            tipo=tipo,
        ))
    if nuevos:
        Convenio.objects.bulk_create(nuevos)
    resumen["convenios"] += len(nuevos)


# --------------------------------------------------------------------------- #
# Entrada pública
# --------------------------------------------------------------------------- #
def ejecutar_importacion(archivo=None, dry_run=False, origen="manual", usuario=None):
    """
    Importa el libro de contratación. Devuelve un dict con el resumen.
    Persiste un ContratoImportLog salvo en dry_run.
    """
    ruta = archivo or settings.CONTRATOS_XLSX_PATH
    wb = _cargar_workbook(ruta)

    resumen = {
        "archivo": str(ruta),
        "filas_leidas": 0,
        "creados": 0,
        "actualizados": 0,
        "sin_match_colaborador": 0,
        "convenios": 0,
        "errores": [],
        "origen": origen,
        "dry_run": dry_run,
    }

    colaboradores = {
        c.cedula: c
        for c in Colaborador.objects.exclude(cedula__isnull=True).exclude(cedula="")
    }

    secop_map = _parse_secop(wb)
    siseg_map = _parse_siseg(wb)

    with transaction.atomic():
        _procesar_hoja_contratos(wb, "Unidad", "unidad", secop_map, siseg_map,
                                 colaboradores, resumen)
        _procesar_hoja_contratos(wb, "Proye. adiciones", "adicion", secop_map, siseg_map,
                                 colaboradores, resumen)
        _procesar_convenios(wb, "Convenios", "convenio", resumen)
        _procesar_convenios(wb, "Otras modalidades", "otra_modalidad", resumen)

        if not dry_run:
            ContratoImportLog.objects.create(
                archivo=resumen["archivo"],
                filas_leidas=resumen["filas_leidas"],
                creados=resumen["creados"],
                actualizados=resumen["actualizados"],
                sin_match_colaborador=resumen["sin_match_colaborador"],
                errores=resumen["errores"],
                origen=origen,
                ejecutado_por=usuario if (usuario and usuario.is_authenticated) else None,
            )
        else:
            transaction.set_rollback(True)

    return resumen
