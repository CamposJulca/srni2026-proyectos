from datetime import date
import json
import tempfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.test.utils import override_settings

from .models import (
    Actividad,
    AlertaProyecto,
    Asignacion,
    Colaborador,
    CuentaCobro,
    Modulo,
    Obligacion,
    Perfil,
    Procedimiento,
    Proyecto,
    ProyectoHistorial,
    ReporteSemanal,
    RiesgoProyecto,
    Rol,
)


class DashboardCruceCobroTests(TestCase):

    def setUp(self):
        # Crear usuario admin para autenticación
        self.user = User.objects.create_user(username='testadmin', password='test1234')
        Perfil.objects.create(user=self.user, rol='admin')
        self.client.login(username='testadmin', password='test1234')

        self.procedimiento = Procedimiento.objects.create(nombre="INSTRUMENTALIZACIÓN")
        self.proyecto = Proyecto.objects.create(nombre="Proyecto A")
        self.modulo = Modulo.objects.create(
            proyecto=self.proyecto,
            nombre="Módulo 1",
            referente="Referente 1",
        )
        self.rol = Rol.objects.create(nombre="Desarrollador")

        self.colaborador = Colaborador.objects.create(
            nombre="Ana Contratista",
            procedimiento=self.procedimiento,
            honorarios=1_000_000,
        )
        Asignacion.objects.create(
            modulo=self.modulo,
            colaborador=self.colaborador,
            rol=self.rol,
        )

        obligacion = Obligacion.objects.create(
            colaborador=self.colaborador,
            descripcion="Cumplir actividades de desarrollo",
        )
        Actividad.objects.create(
            obligacion=obligacion,
            proyecto=self.proyecto,
            actividad_id="A-1",
            descripcion="Actividad completada",
            fecha_inicio=date(2026, 4, 1),
            fecha_fin=date(2026, 4, 10),
            progreso=100,
            estado="completada",
        )
        Actividad.objects.create(
            obligacion=obligacion,
            proyecto=self.proyecto,
            actividad_id="A-2",
            descripcion="Actividad en curso",
            fecha_inicio=date(2026, 4, 11),
            fecha_fin=date(2026, 4, 25),
            progreso=50,
            estado="en_curso",
        )

        CuentaCobro.objects.create(
            colaborador=self.colaborador,
            periodo=date(2026, 4, 7),  # se normaliza al día 1
            numero_cuenta="CC-APR-01",
            valor_cobrado=1_200_000,
            estado="radicada",
        )

    def test_dashboard_data_incluye_cruce_cobro(self):
        res = self.client.get(
            "/api/dashboard/",
            {
                "procedimiento": "INSTRUMENTALIZACIÓN",
                "periodo": "2026-04",
            },
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()

        self.assertIn("cruce_cobro", data)
        self.assertEqual(len(data["cruce_cobro"]), 1)
        fila = data["cruce_cobro"][0]

        self.assertEqual(fila["colaborador"], "Ana Contratista")
        self.assertEqual(fila["actividades_periodo"], 2)
        self.assertEqual(fila["actividades_completadas"], 1)
        self.assertEqual(fila["avance_pct"], 75.0)
        self.assertEqual(fila["cumplimiento_pct"], 50.0)
        self.assertEqual(fila["cobro_pct"], 120.0)
        self.assertEqual(fila["brecha_cumplimiento_cobro"], -70.0)
        self.assertEqual(fila["estado_cruce"], "desalineado")

    def test_dashboard_data_rechaza_periodo_invalido(self):
        res = self.client.get("/api/dashboard/", {"periodo": "2026/04"})
        self.assertEqual(res.status_code, 400)
        self.assertIn("error", res.json())

    def test_gerencial_incluye_semaforo_y_compromisos_por_proyecto(self):
        res = self.client.get("/api/gerencial/")
        self.assertEqual(res.status_code, 200)
        data = res.json()

        self.assertIn("proyectos", data)
        proyecto = data["proyectos"][0]
        self.assertEqual(proyecto["nombre"], "Proyecto A")
        self.assertEqual(proyecto["avance_promedio"], 75.0)
        self.assertEqual(proyecto["personas_asignadas"], 1)
        self.assertEqual(proyecto["modulos_activos"], 1)
        self.assertIn(proyecto["semaforo"], ["verde", "amarillo", "rojo"])
        self.assertEqual(len(data["compromisos_mes"]), 12)

    def test_proyecto_detalle_y_exportacion(self):
        RiesgoProyecto.objects.create(
            proyecto=self.proyecto,
            riesgo="Dependencia externa",
            impacto="alto",
            probabilidad="medio",
        )
        AlertaProyecto.objects.create(
            proyecto=self.proyecto,
            alerta="Actividad próxima a vencer",
            severidad="alta",
        )

        res = self.client.get(f"/api/proyectos/{self.proyecto.pk}/")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["proyecto"]["nombre"], "Proyecto A")
        self.assertEqual(len(data["modulos"]), 1)
        self.assertEqual(len(data["equipo"]), 1)
        self.assertEqual(len(data["cronograma"]), 2)
        self.assertEqual(len(data["riesgos"]), 1)
        self.assertEqual(len(data["alertas"]), 1)

        res = self.client.get(f"/api/proyectos/{self.proyecto.pk}/exportar.csv")
        self.assertEqual(res.status_code, 200)
        self.assertIn("text/csv", res["Content-Type"])


class ArquitecturaCompatibilidadTests(TestCase):

    def test_dashboard_models_reexporta_modelos_partidos(self):
        from dashboard import models
        from dashboard.models.actividades import Actividad as ActividadPartida
        from dashboard.models.organizacion import Colaborador as ColaboradorPartido

        self.assertIs(models.Actividad, ActividadPartida)
        self.assertIs(models.Colaborador, ColaboradorPartido)


class CrudSqlTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username="adminsql", password="test1234")
        Perfil.objects.create(user=self.user, rol="admin")
        self.client.login(username="adminsql", password="test1234")

    def test_sql_query_rechaza_operaciones_no_select(self):
        res = self.client.post(
            "/api/sql/",
            data=json.dumps({"query": "DELETE FROM dashboard_colaborador"}),
            content_type="application/json",
        )

        self.assertEqual(res.status_code, 403)
        self.assertEqual(res.json()["error"], "Solo se permiten consultas SELECT.")

    def test_crud_meta_responde_para_admin(self):
        res = self.client.get("/api/crud/meta/")

        self.assertEqual(res.status_code, 200)
        self.assertIn("colaborador", res.json())
        self.assertIn("proyecto", res.json())
        self.assertIn("riesgo_proyecto", res.json())
        self.assertIn("alerta_proyecto", res.json())


class ProyectoGestionTests(TestCase):

    def test_actividad_completada_sincroniza_progreso(self):
        colaborador = Colaborador.objects.create(nombre="Sincroniza Estado")
        obligacion = Obligacion.objects.create(colaborador=colaborador, descripcion="Obligación")

        actividad = Actividad.objects.create(
            obligacion=obligacion,
            actividad_id="SYNC-1",
            descripcion="Cerrar actividad",
            fecha_inicio=date(2026, 5, 1),
            fecha_fin=date(2026, 5, 2),
            progreso=40,
            estado="completada",
        )

        self.assertEqual(actividad.progreso, 100)

    def test_proyecto_guarda_historial_de_cambios(self):
        proyecto = Proyecto.objects.create(nombre="Proyecto Historial")
        proyecto.estado = "en_riesgo"
        proyecto.prioridad = "alta"
        proyecto.save()

        self.assertEqual(ProyectoHistorial.objects.filter(proyecto=proyecto).count(), 1)
        cambios = proyecto.historial.first().cambios
        self.assertIn("estado", cambios)
        self.assertEqual(cambios["estado"]["despues"], "en_riesgo")


class ReportesTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username="reportero", password="test1234")
        self.colaborador = Colaborador.objects.create(nombre="Reporter Uno")
        Perfil.objects.create(
            user=self.user,
            rol="colaborador",
            colaborador=self.colaborador,
        )
        self.client.login(username="reportero", password="test1234")

    def test_reporte_semanal_guardar_y_listar(self):
        res = self.client.post(
            "/api/reporte-semanal/guardar/",
            data=json.dumps({
                "semana": "ABR S3",
                "que_hizo": "Avancé compromisos",
                "impedimentos": "Ninguno",
            }),
            content_type="application/json",
        )

        self.assertEqual(res.status_code, 200)
        self.assertTrue(res.json()["created"])
        self.assertEqual(ReporteSemanal.objects.count(), 1)

        res = self.client.get("/api/reporte-semanal/", {"semana": "ABR S3"})
        self.assertEqual(res.status_code, 200)
        reportes = res.json()["reportes"]
        self.assertEqual(len(reportes), 1)
        self.assertEqual(reportes[0]["que_hizo"], "Avancé compromisos")

    def test_reporte_semanal_valida_campos_obligatorios(self):
        res = self.client.post(
            "/api/reporte-semanal/guardar/",
            data=json.dumps({"semana": "ABR S3", "que_hizo": ""}),
            content_type="application/json",
        )

        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.json()["error"], "Semana y reporte son obligatorios")


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class EvidenciasTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username="dueno", password="test1234")
        self.otro = User.objects.create_user(username="otro", password="test1234")
        self.colaborador = Colaborador.objects.create(nombre="Dueño Evidencia")
        self.otro_colaborador = Colaborador.objects.create(nombre="Otro Evidencia")
        Perfil.objects.create(
            user=self.user,
            rol="colaborador",
            colaborador=self.colaborador,
        )
        Perfil.objects.create(
            user=self.otro,
            rol="colaborador",
            colaborador=self.otro_colaborador,
        )
        obligacion = Obligacion.objects.create(
            colaborador=self.colaborador,
            descripcion="Adjuntar evidencia",
        )
        self.actividad = Actividad.objects.create(
            obligacion=obligacion,
            actividad_id="EV-1",
            descripcion="Actividad con evidencia",
            fecha_inicio=date(2026, 4, 1),
            fecha_fin=date(2026, 4, 2),
        )

    def test_dueno_sube_y_lista_evidencia(self):
        self.client.login(username="dueno", password="test1234")
        archivo = SimpleUploadedFile("evidencia.txt", b"contenido", content_type="text/plain")

        res = self.client.post(
            f"/api/evidencias/{self.actividad.pk}/subir/",
            data={"archivo": archivo, "comentario": "Entregado"},
        )

        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["nombre"], "evidencia.txt")

        res = self.client.get(f"/api/evidencias/{self.actividad.pk}/")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["total"], 1)
        self.assertEqual(res.json()["evidencias"][0]["comentario"], "Entregado")

    def test_otro_usuario_no_puede_subir_evidencia(self):
        self.client.login(username="otro", password="test1234")
        archivo = SimpleUploadedFile("evidencia.txt", b"contenido", content_type="text/plain")

        res = self.client.post(
            f"/api/evidencias/{self.actividad.pk}/subir/",
            data={"archivo": archivo, "comentario": "No autorizado"},
        )

        self.assertEqual(res.status_code, 403)
        self.assertEqual(res.json()["error"], "No autorizado")


class ImportacionTests(TestCase):

    def setUp(self):
        self.user = User.objects.create_user(username="adminimport", password="test1234")
        Perfil.objects.create(user=self.user, rol="admin")
        self.client.login(username="adminimport", password="test1234")

    def test_carga_post_sin_archivo_retorna_400(self):
        res = self.client.post("/carga/")

        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.json()["error"], "No se recibió ningún archivo.")


def _crear_workbook_contratos(ruta):
    """Genera un .xlsx mínimo con la estructura del libro maestro (valores literales)."""
    import openpyxl

    wb = openpyxl.Workbook()

    # --- Unidad ---
    unidad = wb.active
    unidad.title = "Unidad"
    unidad.append([
        "ESTADO", "CEDULA", "NOMBRE CONTRATISTA", "EQUIPO", "ID", "CONTRATO",
        "VALOR HONORARIOS", "VALOR REAL CONTRATO SEGÚN FECHA DE INICIO",
        "VALOR TOTAL CANCELADO", "SALDO CONTRATO", "%  EJEC.", "IBC",
        "PAGO 1", "PAGO 2", "PAGO 3", "PAGO 4", "PAGO 5",
    ])
    unidad.append([
        "EJECUCIÓN", 111, "ANA TEST", "AIDI", 1, "1-2026",
        1000, 10000, 5000, 5000, 0.5, 100,
        1000, 1000, 1000, 1000, 1000,
    ])
    unidad.append([
        "EJECUCIÓN", 222, "LUIS TEST", "GIS", 2, "2-2026",
        2000, 20000, 4000, 16000, 0.2, 200,
        2000, 2000, None, None, None,
    ])

    # --- SECOP II ---
    secop = wb.create_sheet("SECOP II")
    secop.append([
        "ESTADO", "CONTRATO", "CEDULA", "NOMBRE CONTRATISTA",
        "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
        "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
    ])
    secop.append(["EJECUCIÓN", "1-2026", 111, "ANA TEST",
                  "Pagado", "Pagado", "Pagado", "Pagado", "Pagado",
                  None, None, None, None, None, None, None])
    secop.append(["EJECUCIÓN", "2-2026", 222, "LUIS TEST",
                  "Pagado", "Pagado", None, None, None,
                  None, None, None, None, None, None, None])

    # --- SISEG (encabezados en la fila 2) ---
    siseg = wb.create_sheet("SISEG")
    siseg.append(["INFORMACIÓN CONTRATO"])  # fila 1: agrupación
    siseg.append([
        "ESTADO", "CONTRATO", "CEDULA", "NOMBRE CONTRATISTA", "EQUIPO",
        "INICIO", "TERMINACIÓN",
        "CÓDIGO", "APROBADO", "SS", "CÓDIGO", "APROBADO", "SS",
        "CÓDIGO", "APROBADO", "SS", "CÓDIGO", "APROBADO", "SS",
        "CÓDIGO", "APROBADO", "SS",
    ])
    siseg.append([
        "EJECUCIÓN", "1-2026", 111, "ANA TEST", "AIDI", None, None,
        "COD1", None, "ENERO", "COD2", None, "FEBRERO",
        "COD3", None, "MARZO", "COD4", None, "ABRIL",
    ])
    siseg.append([
        "EJECUCIÓN", "2-2026", 222, "LUIS TEST", "GIS", None, None,
        "COD1", None, "ENERO",
    ])

    # --- Convenios ---
    conv = wb.create_sheet("Convenios")
    conv.append([
        "ENTIDAD", "ARTICULADOR", "SEYCO", "No. CONVENIO", "FECHA SUSCRIPCIÓN",
        "FECHA TERMINACIÓN", "No. MODIFICATORIOS", "CONTRATISTA",
        "CLASE DE CONTRATO", "VALOR INICIAL CONTRATO", "OBJETO",
        "LINK DEL ACUERDO O CONVENIO ",
    ])
    conv.append([
        "Entidad X", "Art Y", "NI-0001", 99, None,
        "Renovacion automatica", None, "Contratista Z",
        "Convenio Interadministrativo", 0, "Objeto de prueba", "http://x",
    ])

    wb.save(ruta)


class ContratacionTests(TestCase):

    def setUp(self):
        self.tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        self.tmp.close()
        _crear_workbook_contratos(self.tmp.name)

        self.user = User.objects.create_user(username="admincontrato", password="test1234")
        Perfil.objects.create(user=self.user, rol="admin")
        self.client.login(username="admincontrato", password="test1234")

    def _importar(self):
        from dashboard.domains.contratacion.importador import ejecutar_importacion
        return ejecutar_importacion(archivo=self.tmp.name, origen="manual", usuario=self.user)

    def test_clasificar_cruce_unit(self):
        from dashboard.domains.contratacion.utils import clasificar_cruce_fisico_financiero
        self.assertEqual(clasificar_cruce_fisico_financiero(None, 50), "sin_datos")
        self.assertEqual(clasificar_cruce_fisico_financiero(48, 50), "alineado")
        self.assertEqual(clasificar_cruce_fisico_financiero(35, 50), "alerta")
        self.assertEqual(clasificar_cruce_fisico_financiero(10, 50), "desalineado")

    def test_importador_parsea_y_es_idempotente(self):
        from dashboard.models import (
            Contrato, ContratoImportLog, Convenio, PagoContrato, SeguridadSocialContrato,
        )

        r1 = self._importar()
        self.assertEqual(r1["creados"], 2)
        self.assertEqual(r1["actualizados"], 0)
        self.assertEqual(r1["convenios"], 1)
        self.assertEqual(r1["errores"], [])

        c1 = Contrato.objects.get(numero_contrato="1-2026", modalidad="unidad")
        self.assertEqual(c1.nombre_contratista, "ANA TEST")
        self.assertEqual(c1.cedula, "111")
        self.assertEqual(float(c1.valor_total_contrato), 10000)
        self.assertEqual(float(c1.valor_cancelado), 5000)
        self.assertEqual(float(c1.porcentaje_ejecucion), 50.0)  # 0.5 -> 50%
        self.assertEqual(c1.pagos.count(), 5)
        self.assertEqual(c1.pagos.filter(estado_secop="pagado").count(), 5)
        self.assertEqual(c1.seguridad_social.count(), 4)
        self.assertEqual(Convenio.objects.count(), 1)

        # Segunda corrida: idempotente (0 nuevos, mismos totales).
        r2 = self._importar()
        self.assertEqual(r2["creados"], 0)
        self.assertEqual(r2["actualizados"], 2)
        self.assertEqual(Contrato.objects.count(), 2)
        self.assertEqual(PagoContrato.objects.count(), 7)
        self.assertEqual(SeguridadSocialContrato.objects.count(), 5)
        self.assertEqual(Convenio.objects.count(), 1)
        self.assertEqual(ContratoImportLog.objects.count(), 2)

    def test_dry_run_no_escribe(self):
        from dashboard.models import Contrato, ContratoImportLog
        from dashboard.domains.contratacion.importador import ejecutar_importacion

        resumen = ejecutar_importacion(archivo=self.tmp.name, dry_run=True)
        self.assertEqual(resumen["creados"], 2)
        self.assertEqual(Contrato.objects.count(), 0)
        self.assertEqual(ContratoImportLog.objects.count(), 0)

    def test_payload_kpis_y_cruce(self):
        # Colaborador con cédula que empata con el contrato 1-2026 y avance físico 20%.
        colaborador = Colaborador.objects.create(nombre="Ana Fisica", cedula="111", honorarios=1000)
        proyecto = Proyecto.objects.create(nombre="P Contratos")
        obligacion = Obligacion.objects.create(colaborador=colaborador, descripcion="obl")
        Actividad.objects.create(
            obligacion=obligacion, proyecto=proyecto, actividad_id="AC-1",
            descripcion="a", fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 1, 5),
            progreso=20, estado="en_curso",
        )
        self._importar()
        # Enlaza el colaborador creado después de importar.
        from dashboard.models import Contrato
        Contrato.objects.filter(cedula="111").update(colaborador=colaborador)

        res = self.client.get("/api/gerencial/contratos/")
        self.assertEqual(res.status_code, 200)
        data = res.json()

        self.assertEqual(data["kpis"]["total_contratos"], 2)
        self.assertEqual(data["kpis"]["valor_total"], 30000.0)
        self.assertEqual(data["kpis"]["valor_cancelado"], 9000.0)

        fila = next(f for f in data["cruce_fisico_financiero"] if f["numero_contrato"] == "1-2026")
        self.assertEqual(fila["avance_fisico"], 20.0)
        self.assertEqual(fila["avance_financiero"], 50.0)
        self.assertEqual(fila["estado_cruce"], "desalineado")

    def tearDown(self):
        import os
        os.unlink(self.tmp.name)
